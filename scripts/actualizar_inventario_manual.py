# -*- coding: utf-8 -*-
"""Actualiza el inventario a la NUEVA BASE real (conteo físico manual en planta).

Fuente: 'C:\\Users\\andre\\OneDrive\\Escritorio\\DELCA\\DELCA INVENTARIO. MIGRAR.xlsx'
  - Hoja 'Materia prima ': productos de materia prima COMPLETOS (código y costo).
  - Hoja 'Consumible'    : consumibles; código y costo VACÍOS (se completan
                           manualmente en el sistema después de la migración).

Comportamiento:
  - UPSERT por SKU (código para MP; SKU generado nombre+dimensión para consumibles).
  - El stock/stock_kg se FIJAN al valor contado (nueva base) mediante un
    movimiento AJUSTE en el Kardex (no suma: ajusta al conteo real).
  - Los consumibles quedan con costo 0 y código pendiente (los completa el usuario).
  - Los productos existentes que NO están en el Excel no se tocan (se reportan).

Uso: python scripts/actualizar_inventario_manual.py [--dry-run | --ejecutar]
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import sqlite3
import sys
import unicodedata
from datetime import datetime

import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BASE_DIR, "delca.db")
BACKUP_DIR = os.path.join(BASE_DIR, "backups", "migracion")
EXCEL_PATH = os.environ.get(
    "DELCA_ARCHIVO_INVENTARIO",
    r"C:\Users\andre\OneDrive\Escritorio\DELCA\DELCA INVENTARIO. MIGRAR.xlsx",
)
HOJA_MP = "Materia prima "
HOJA_CONS = "Consumible"

UNIDADES_VALIDAS = {"UNIDAD", "KG", "LT", "MT", "CAJA", "PAQ", "ROLLO"}


def normalizar_sku(texto: str) -> str:
    """Genera un SKU legible desde texto (sin espacios/acentos/símbolos)."""
    s = unicodedata.normalize("NFD", texto)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^0-9A-Za-z]", "", s.upper())
    return s[:50]


def unidad_normalizada(u: str | None, por_defecto: str = "UNIDAD") -> str:
    u = (u or "").strip().upper()
    for valida in UNIDADES_VALIDAS:
        if valida in u or u in valida:
            return valida
    return por_defecto


def num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def leer_mp(ws) -> list[dict]:
    """Hoja Materia prima: Nombre, Codigo, Costo, Kg, Rollos, Minimo, Unidad, Fecha."""
    filas = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        nombre = str(row[0]).strip() if row[0] is not None else ""
        if not nombre:
            continue
        filas.append({
            "fila": i, "nombre": nombre, "codigo": str(row[1]).strip() if row[1] is not None else "",
            "costo": num(row[2]), "kg": num(row[3]), "rollos": num(row[4]),
            "minimo": num(row[5]), "unidad": unidad_normalizada(row[6], "ROLLO"),
            "fecha": row[7],
        })
    return filas


def leer_consumibles(ws) -> list[dict]:
    """Hoja Consumible: NOMBRE, Dimension, Codigo(vacío), Cant x Caja, Und por caja,
    Unidad DE SALIDA, Costo(vacío), Inventario Actual, Minimo, Fecha."""
    filas = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        nombre = str(row[0]).strip() if row[0] is not None else ""
        if not nombre:
            continue
        dimension = str(row[1]).strip() if row[1] is not None else ""
        inventario = num(row[7])
        cant_caja = num(row[3])
        nombre_final = f"{nombre} {dimension}".strip() if dimension else nombre
        filas.append({
            "fila": i, "nombre": nombre_final, "nombre_base": nombre,
            "dimension": dimension, "codigo": str(row[2]).strip() if row[2] is not None else "",
            "cant_caja": cant_caja, "costo": num(row[6]),
            "inventario": inventario, "minimo": num(row[8]),
            "unidad": unidad_normalizada(row[5], "CAJA"),
            "fecha": row[9],
            "kg": cant_caja * inventario if cant_caja > 0 else 0.0,
        })
    return filas


def main() -> None:
    parser = argparse.ArgumentParser(description="Actualizar inventario a la nueva base (conteo físico)")
    grupo = parser.add_mutually_exclusive_group()
    grupo.add_argument("--dry-run", action="store_true", help="Solo simula (por defecto)")
    grupo.add_argument("--ejecutar", action="store_true", help="Aplica con backup")
    args = parser.parse_args()
    ejecutar = args.ejecutar

    if sys.stdout and hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    if not os.path.exists(EXCEL_PATH):
        print(f"[ERROR] No se encuentra el Excel: {EXCEL_PATH}")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    mp = leer_mp(wb[HOJA_MP])
    cons = leer_consumibles(wb[HOJA_CONS])
    print(f"[EXCEL] Materia prima: {len(mp)} filas | Consumibles: {len(cons)} filas")

    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    cur = con.cursor()

    plan = []  # (sku, categoria, nombre, costo, stock, stock_kg, minimo, unidad, fecha, fila, codigo)
    for f in mp:
        sku = f["codigo"] or normalizar_sku(f["nombre"])
        plan.append({
            "sku": sku, "categoria": "MATERIA_PRIMA", "nombre": f["nombre"],
            "costo": f["costo"], "stock": f["rollos"], "stock_kg": f["kg"],
            "minimo": f["minimo"], "unidad": f["unidad"], "fecha": f["fecha"],
            "fila": f["fila"], "codigo": f["codigo"], "pendiente": False,
        })
    for f in cons:
        sku = f["codigo"] or normalizar_sku(f["nombre"])
        plan.append({
            "sku": sku, "categoria": "CONSUMIBLE", "nombre": f["nombre"],
            "costo": f["costo"], "stock": f["inventario"], "stock_kg": f["kg"],
            "minimo": f["minimo"], "unidad": f["unidad"], "fecha": f["fecha"],
            "fila": f["fila"], "codigo": f["codigo"], "pendiente": (not f["codigo"] and f["costo"] <= 0),
        })

    # Productos existentes no listados (reporte)
    skus_plan = {p["sku"] for p in plan}
    existentes_no_listados = [
        (r["sku"], r["nombre"], r["categoria"])
        for r in cur.execute("SELECT sku, nombre, categoria FROM productos WHERE sku NOT IN ({})".format(
            ",".join("?" * len(skus_plan))), tuple(skus_plan)).fetchall()
    ] if skus_plan else []

    n_crear = n_act = n_ajuste = 0
    pendientes = []
    detalles = []
    for p in plan:
        existente = cur.execute("SELECT id, stock, stock_kg FROM productos WHERE sku=?", (p["sku"],)).fetchone()
        if existente:
            n_act += 1
            dif = p["stock"] - float(existente["stock"])
            dif_kg = p["stock_kg"] - float(existente["stock_kg"])
        else:
            n_crear += 1
            dif = p["stock"]
            dif_kg = p["stock_kg"]
        if abs(dif) > 0 or abs(dif_kg) > 0:
            n_ajuste += 1
        if p["pendiente"]:
            pendientes.append(p["nombre"])
        detalles.append({
            "sku": p["sku"], "nombre": p["nombre"], "categoria": p["categoria"],
            "crear": existente is None, "costo": p["costo"], "stock": p["stock"],
            "stock_kg": p["stock_kg"], "minimo": p["minimo"], "unidad": p["unidad"],
            "dif": dif, "dif_kg": dif_kg, "fecha": p["fecha"], "pendiente": p["pendiente"],
        })

    print(f"[PLAN] crear: {n_crear} | actualizar: {n_act} | ajustes de stock: {n_ajuste}")
    print(f"[PENDIENTES] consumibles con codigo/costo por completar: {len(pendientes)}")
    print(f"[NO LISTADOS] productos existentes que NO estan en el Excel: {len(existentes_no_listados)}")
    for e in existentes_no_listados[:10]:
        print(f"    {e[0]} | {e[1]} | {e[2]}")

    print("\n=== Muestra (primeras 8) ===")
    for d in detalles[:8]:
        accion = "CREAR" if d["crear"] else "ACTUALIZAR"
        print("  [{}] {} | sku={} | cat={} | stock={} (dif {}) | kg={} (dif {}) | costo={:,.0f} | und={} {}".format(
            accion, d["nombre"], d["sku"], d["categoria"], d["stock"], d["dif"],
            d["stock_kg"], d["dif_kg"], d["costo"], d["unidad"], "(pendiente cod/costo)" if d["pendiente"] else ""))

    if not ejecutar:
        print("\n[DRY-RUN] No se escribio nada. Usa --ejecutar para aplicar.")
        return

    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = os.path.join(BACKUP_DIR, f"delca_pre_inventario_manual_{ts}.db")
    shutil.copy2(DB_PATH, backup)
    print(f"\n[BACKUP] {backup}")

    con.execute("BEGIN")
    fecha_ajuste = datetime.now()
    for d in detalles:
        if d["crear"]:
            cur.execute(
                "INSERT INTO productos (nombre, sku, descripcion, categoria, stock, stock_kg, "
                "stock_minimo, costo_unitario, precio_venta, unidad_medida, activo) "
                "VALUES (?, ?, NULL, ?, ?, ?, ?, ?, 0, ?, 1)",
                (d["nombre"], d["sku"], d["categoria"], d["stock"], d["stock_kg"],
                 d["minimo"], d["costo"], d["unidad"]),
            )
            pid = cur.lastrowid
        else:
            cur.execute(
                "UPDATE productos SET nombre=?, categoria=?, stock=?, stock_kg=?, "
                "stock_minimo=?, costo_unitario=?, unidad_medida=?, activo=1 WHERE sku=?",
                (d["nombre"], d["categoria"], d["stock"], d["stock_kg"],
                 d["minimo"], d["costo"], d["unidad"], d["sku"]),
            )
            pid = cur.execute("SELECT id FROM productos WHERE sku=?", (d["sku"],)).fetchone()[0]
        # Ajuste de inventario (solo si hay diferencia)
        if abs(d["dif"]) > 0 or abs(d["dif_kg"]) > 0:
            fecha = d["fecha"] if d["fecha"] else fecha_ajuste
            cur.execute(
                "INSERT INTO movimientos_inventario (producto_id, tipo, cantidad, costo_unitario, "
                "referencia, observaciones, fecha, cantidad_kg) "
                "VALUES (?, 'AJUSTE', ?, ?, 'INVENTARIO_MANUAL', "
                "'Inventario manual (nueva base) - conteo en planta', ?, ?)",
                (pid, d["dif"], d["costo"], fecha, d["dif_kg"]),
            )
    con.commit()

    print(f"\n=== REPORTE FINAL ===")
    print(f"  Creados: {n_crear} | Actualizados: {n_act} | Ajustes de stock: {n_ajuste}")
    print(f"  Consumibles con codigo/costo pendientes: {len(pendientes)} (los completa el usuario)")
    print(f"  Productos existentes NO listados (sin tocar): {len(existentes_no_listados)}")
    print(f"  Backup: {backup}")
    con.close()
    print("ACTUALIZACION APLICADA OK")


if __name__ == "__main__":
    main()