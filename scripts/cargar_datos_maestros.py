"""
Cargador de datos maestros desde Excel a la base de datos.

Lee el archivo data/datos_maestros.xlsx y pobla/actualiza las
tablas del sistema con la informacion de referencia.

Uso:
    python scripts/cargar_datos_maestros.py

Para actualizar precios: modifique el Excel y re-ejecute.
Los registros existentes se actualizan (UPSERT por nombre/codigo).
"""

import sys
from decimal import Decimal
from pathlib import Path

# ── Asegurar que el proyecto esta en el path ────────────────────────
BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

from openpyxl import load_workbook

# ── Registrar modelos ANTES de cualquier operacion de BD ────────────
import src.database.registry  # noqa: F401

from src.database.engine import DB_PATH, engine, get_session
from src.database.base import Base
from src.modules.llantas.models.marca_llanta_model import MarcaLlanta
from src.modules.llantas.models.dimension_llanta_model import DimensionLlanta
from src.modules.llantas.models.diseno_llanta_model import DisenoLlanta, TIPOS_DISENO
from src.modules.llantas.models.causa_rechazo_model import CausaRechazo
from src.modules.inventario.models.producto_model import Producto
from src.modules.inventario.models.precio_producto_model import PrecioProducto
from src.modules.inventario.models.inventario_config_models import (
    RecetaProduccion,
)

DATA_FILE = BASE_DIR / "data" / "datos_maestros.xlsx"


# ======================================================================
#  Helpers
# ======================================================================


def _buscar_dimension_por_display(session, medida_display: str) -> DimensionLlanta | None:
    """Busca una dimensión por su display canónico (ej. '205/55 R16').

    Normaliza el texto eliminando espacios para comparar contra el
    display calculado del modelo.
    """
    normalizado = medida_display.replace(" ", "").upper()
    for dim in session.query(DimensionLlanta).all():
        if dim.display.replace(" ", "").upper() == normalizado:
            return dim
    return None


# ======================================================================
#  Cargadores por hoja
# ======================================================================


def _cargar_marcas(ws, session):
    """Hoja: Marcas | columnas: nombre"""
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        nombre = (row[0] or "").strip()
        if not nombre:
            continue
        existente = session.query(MarcaLlanta).filter_by(nombre=nombre).first()
        if existente:
            continue  # ya existe, no duplica
        session.add(MarcaLlanta(nombre=nombre))
        count += 1
    session.commit()
    return count


def _cargar_medidas(ws, session):
    """Hoja: Medidas | columnas: ancho, perfil, rin"""
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        ancho, perfil, rin = row[0], row[1], row[2]
        if ancho is None or perfil is None or rin is None:
            continue
        existente = (
            session.query(DimensionLlanta)
            .filter_by(ancho=int(ancho), perfil=int(perfil), rin=int(rin))
            .first()
        )
        if existente:
            continue
        session.add(DimensionLlanta(ancho=int(ancho), perfil=int(perfil), rin=int(rin)))
        count += 1
    session.commit()
    return count


def _cargar_disenos(ws, session):
    """Hoja: Disenos | columnas: nombre, tipo (opcional; por defecto MIXTO).

    Desde v2.2.0 los diseños son independientes de la marca: la columna
    legacy 'marca_nombre' ya no se usa. El 'tipo' acepta MIXTO/TRACCION/
    DIRECCIONAL (default MIXTO).
    """
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        nombre = (row[0] or "").strip()
        if not nombre or nombre.startswith("'"):
            continue
        # Si el nombre ya viene con sufijo de marca legacy ("D-100 Goodyear"),
        # solo interesa la parte del diseño.
        nombre_limpio = nombre.split(" ")[0].strip()
        tipo = (row[1] or "MIXTO").strip().upper()
        if tipo not in TIPOS_DISENO:
            tipo = "MIXTO"
        existente = (
            session.query(DisenoLlanta)
            .filter_by(nombre=nombre_limpio)
            .first()
        )
        if existente:
            continue
        session.add(DisenoLlanta(nombre=nombre_limpio, tipo=tipo))
        count += 1
    session.commit()
    return count


def _cargar_productos(ws, session):
    """Hoja: Productos | columnas: nombre, sku, descripcion, categoria, costo_unitario, unidad_medida, stock_minimo"""
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        nombre = (row[0] or "").strip()
        sku = (row[1] or "").strip()
        if not nombre or not sku:
            continue
        descripcion = (row[2] or "").strip()
        categoria = (row[3] or "").strip()
        costo = float(row[4]) if row[4] is not None else 0
        unidad = (row[5] or "UNidad").strip()
        stock_min = float(row[6]) if row[6] is not None else 0

        existente = session.query(Producto).filter_by(sku=sku).first()
        if existente:
            # Actualizar datos (precios cambian en el tiempo)
            existente.nombre = nombre
            existente.descripcion = descripcion or existente.descripcion
            existente.categoria = categoria or existente.categoria
            existente.costo_unitario = costo
            existente.unidad_medida = unidad
            existente.stock_minimo = stock_min
            count += 1
            continue
        session.add(Producto(
            nombre=nombre, sku=sku, descripcion=descripcion,
            categoria=categoria, costo_unitario=costo,
            unidad_medida=unidad, stock_minimo=stock_min,
        ))
        count += 1
    session.commit()
    return count


def _cargar_precios(ws, session):
    """Hoja: PreciosProducto | columnas: medida_display, diseno_nombre, costo_fabricacion, precio_minimo, precio_medio, precio_normal"""
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        medida_display = (row[0] or "").strip()
        diseno_nombre = (row[1] or "").strip()
        if not medida_display or not diseno_nombre:
            continue

        # Buscar medida por display "ancho/perfil Rrin" (ej. "205/55 R16")
        medida = _buscar_dimension_por_display(session, medida_display)
        if not medida:
            print(f"  [AVISO] Medida no encontrada: {medida_display}")
            continue

        # Desde v2.2.0 los diseños son independientes: nombre simple "D-100"
        diseno_nombre_limpio = diseno_nombre.split(" ")[0].strip()
        diseno = session.query(DisenoLlanta).filter_by(nombre=diseno_nombre_limpio).first()
        if not diseno:
            print(f"  [AVISO] Diseno no encontrado: {diseno_nombre_limpio}")
            continue

        existente = (
            session.query(PrecioProducto)
            .filter_by(dimension_id=medida.id, diseno_id=diseno.id)
            .first()
        )
        vals = {
            "costo_fabricacion": Decimal(str(row[2])) if row[2] is not None else Decimal("0"),
            "precio_minimo": Decimal(str(row[3])) if row[3] is not None else Decimal("0"),
            "precio_medio": Decimal(str(row[4])) if row[4] is not None else Decimal("0"),
            "precio_normal": Decimal(str(row[5])) if row[5] is not None else Decimal("0"),
        }
        if existente:
            for k, v in vals.items():
                setattr(existente, k, v)
        else:
            session.add(PrecioProducto(
                dimension_id=medida.id, diseno_id=diseno.id, **vals
            ))
        count += 1
    session.commit()
    return count


def _cargar_causas_rechazo(ws, session):
    """Hoja: CausasRechazo | columnas: codigo, descripcion, categoria"""
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        codigo = (row[0] or "").strip()
        descripcion = (row[1] or "").strip()
        if not codigo or not descripcion:
            continue
        categoria = (row[2] or "").strip() or None
        existente = session.query(CausaRechazo).filter_by(codigo=codigo).first()
        if existente:
            existente.descripcion = descripcion
            existente.categoria = categoria
        else:
            session.add(CausaRechazo(
                codigo=codigo, descripcion=descripcion, categoria=categoria
            ))
        count += 1
    session.commit()
    return count


def _cargar_recetas(ws, session):
    """Hoja: RecetasProduccion | columnas: medida_display, diseno_nombre, producto_sku, cantidad, unidad"""
    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        medida_display = (row[0] or "").strip()
        diseno_nombre = (row[1] or "").strip()
        producto_sku = (row[2] or "").strip()
        if not medida_display or not diseno_nombre or not producto_sku:
            continue

        medida = _buscar_dimension_por_display(session, medida_display)
        if not medida:
            continue
        diseno_nombre_limpio = diseno_nombre.split(" ")[0].strip()
        diseno = session.query(DisenoLlanta).filter_by(nombre=diseno_nombre_limpio).first()
        if not diseno:
            continue
        producto = session.query(Producto).filter_by(sku=producto_sku).first()
        if not producto:
            print(f"  [AVISO] Producto SKU '{producto_sku}' no encontrado")
            continue

        cantidad = float(row[3]) if row[3] is not None else 0
        unidad = (row[4] or "UNidad").strip()

        existente = (
            session.query(RecetaProduccion)
            .filter_by(
                diseno_id=diseno.id, dimension_id=medida.id, producto_id=producto.id
            )
            .first()
        )
        if existente:
            existente.cantidad = cantidad
            existente.unidad = unidad
        else:
            session.add(RecetaProduccion(
                diseno_id=diseno.id, dimension_id=medida.id,
                producto_id=producto.id, cantidad=cantidad, unidad=unidad,
            ))
        count += 1
    session.commit()
    return count


# ======================================================================
#  Main
# ======================================================================


def inicializar_datos_maestros(only_if_empty: bool = True) -> int:
    """Carga datos maestros desde el Excel (idempotente).

    Pensada para el arranque del programa (main.py): en una instalación
    limpia (BD nueva sin catálogos), puebla marcas/medidas/diseños/
    productos/precios/causas/recetas automáticamente.

    Args:
        only_if_empty: si True, solo carga cuando no hay marcas registradas
            (evita duplicar en instalaciones ya pobladas).

    Returns:
        Número de registros procesados (0 si no aplica).
    """
    if not DATA_FILE.exists():
        print(f"[datos_maestros] No se encuentra {DATA_FILE} — saltando")
        return 0

    # Idempotencia: si ya hay catálogos cargados, no re-procesar
    if only_if_empty:
        from src.modules.llantas.models.marca_llanta_model import MarcaLlanta

        with get_session() as session:
            if session.query(MarcaLlanta).count() > 0:
                print("[datos_maestros] Catálogos ya poblados — saltando")
                return 0

    print(f"[datos_maestros] Leyendo: {DATA_FILE}")
    wb = load_workbook(DATA_FILE, data_only=True)
    Base.metadata.create_all(bind=engine)

    total = {}
    with get_session() as session:
        if "Marcas" in wb.sheetnames:
            total["Marcas"] = _cargar_marcas(wb["Marcas"], session)
        if "Medidas" in wb.sheetnames:
            total["Medidas"] = _cargar_medidas(wb["Medidas"], session)
        if "Disenos" in wb.sheetnames:
            total["Disenos"] = _cargar_disenos(wb["Disenos"], session)
        if "Productos" in wb.sheetnames:
            total["Productos"] = _cargar_productos(wb["Productos"], session)
        if "PreciosProducto" in wb.sheetnames:
            total["PreciosProducto"] = _cargar_precios(wb["PreciosProducto"], session)
        if "CausasRechazo" in wb.sheetnames:
            total["CausasRechazo"] = _cargar_causas_rechazo(wb["CausasRechazo"], session)
        if "RecetasProduccion" in wb.sheetnames:
            total["RecetasProduccion"] = _cargar_recetas(wb["RecetasProduccion"], session)

    print("[datos_maestros] Resumen:", total)
    return sum(total.values())


def main():
    if not DATA_FILE.exists():
        print(f"[ERROR] No se encuentra el archivo: {DATA_FILE}")
        print("Ejecute primero: python scripts/crear_formato_datos_maestros.py")
        sys.exit(1)

    print(f"Leyendo: {DATA_FILE}")
    print(f"Base de datos: {DB_PATH}")
    print()

    wb = load_workbook(DATA_FILE, data_only=True)

    # Crear tablas si no existen (incluye causas_rechazo)
    Base.metadata.create_all(bind=engine)

    total = {}

    with get_session() as session:
        if "Marcas" in wb.sheetnames:
            n = _cargar_marcas(wb["Marcas"], session)
            total["Marcas"] = n
            print(f"  Marcas: {n} registros procesados")

        if "Medidas" in wb.sheetnames:
            n = _cargar_medidas(wb["Medidas"], session)
            total["Medidas"] = n
            print(f"  Medidas: {n} registros procesados")

        if "Disenos" in wb.sheetnames:
            n = _cargar_disenos(wb["Disenos"], session)
            total["Disenos"] = n
            print(f"  Disenos: {n} registros procesados")

        if "Productos" in wb.sheetnames:
            n = _cargar_productos(wb["Productos"], session)
            total["Productos"] = n
            print(f"  Productos: {n} registros procesados")

        if "PreciosProducto" in wb.sheetnames:
            n = _cargar_precios(wb["PreciosProducto"], session)
            total["PreciosProducto"] = n
            print(f"  PreciosProducto: {n} registros procesados")

        if "CausasRechazo" in wb.sheetnames:
            n = _cargar_causas_rechazo(wb["CausasRechazo"], session)
            total["CausasRechazo"] = n
            print(f"  CausasRechazo: {n} registros procesados")

        if "RecetasProduccion" in wb.sheetnames:
            n = _cargar_recetas(wb["RecetasProduccion"], session)
            total["RecetasProduccion"] = n
            print(f"  RecetasProduccion: {n} registros procesados")

    print()
    print("--- Resumen ---")
    for tabla, n in total.items():
        print(f"  {tabla}: {n}")
    print()
    print("[OK] Carga completada.")


if __name__ == "__main__":
    main()
