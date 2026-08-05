"""
Genera la plantilla Excel "datos_maestros.xlsx" con todas las hojas
de datos de referencia que el sistema necesita.

Ejecutar: python scripts/crear_formato_datos_maestros.py
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
OUTPUT = DATA_DIR / "datos_maestros.xlsx"

HEADER_FILL = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
NOTE_FONT = Font(name="Calibri", italic=True, color="888888", size=9)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _estilizar_header(ws, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = 25
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _agregar_nota(ws, fila, col, texto):
    cell = ws.cell(row=fila, column=col, value=texto)
    cell.font = NOTE_FONT


def build():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # ── SHEET 1: Marcas ─────────────────────────────────────────────
    ws = wb.active
    ws.title = "Marcas"
    _estilizar_header(ws, ["nombre"])
    ws.append(["Goodyear"])
    ws.append(["Michelin"])
    ws.append(["Bridgestone"])
    _agregar_nota(ws, 6, 1, "Agregue tantas filas como necesite. 'nombre' debe ser único.")

    # ── SHEET 2: Medidas ────────────────────────────────────────────
    ws = wb.create_sheet("Medidas")
    _estilizar_header(ws, ["ancho", "perfil", "rin"])
    ws.append([205, 55, 16])
    ws.append([225, 60, 18])
    ws.append([195, 65, 15])
    _agregar_nota(ws, 6, 1, "ancho(mm), perfil, rin(pulgadas). Ej: 205/55 R16")

    # ── SHEET 3: Diseños ────────────────────────────────────────────
    ws = wb.create_sheet("Disenos")
    _estilizar_header(ws, ["nombre", "marca_nombre"])
    ws.append(["D-100", "Goodyear"])
    ws.append(["D-200", "Michelin"])
    _agregar_nota(ws, 5, 1, "'marca_nombre' debe coincidir con un nombre en la hoja Marcas.")

    # ── SHEET 4: Productos (Materia prima y consumibles) ────────────
    ws = wb.create_sheet("Productos")
    _estilizar_header(ws, [
        "nombre", "sku", "descripcion", "categoria",
        "costo_unitario", "unidad_medida", "stock_minimo",
    ])
    ws.append(["Caucho Base", "MP-CAUCHO", "Caucho natural para reencauche",
               "Materia Prima", 4500.00, "KG", 100])
    ws.append(["Cemento Vulcanizante", "MP-CEMENTO", "Cemento para unión banda-carcasa",
               "Materia Prima", 3200.00, "KG", 50])
    ws.append(["Banda de Rodadura", "MP-BANDA", "Banda de rodadura precurada",
               "Consumible", 8500.00, "UNidad", 20])
    _agregar_nota(ws, 7, 1, "unidad_medida: UNidad, KG, LT, MT, CAJA, PAQ")

    # ── SHEET 5: PreciosProducto ────────────────────────────────────
    ws = wb.create_sheet("PreciosProducto")
    _estilizar_header(ws, [
        "medida_display", "diseno_nombre",
        "costo_fabricacion", "precio_minimo", "precio_medio", "precio_normal",
    ])
    ws.append(["205/55 R16", "D-100", 35000, 45000, 52000, 58000])
    ws.append(["225/60 R18", "D-200", 42000, 55000, 62000, 70000])
    _agregar_nota(ws, 6, 1,
                  "medida_display: formato 'ancho/perfil Rrin'. "
                  "diseno_nombre debe coincidir con hoja Disenos. "
                  "Precios se actualizarán en el tiempo — modifique el Excel y re-ejecute el cargador.")

    # ── SHEET 6: CausasRechazo ─────────────────────────────────────
    ws = wb.create_sheet("CausasRechazo")
    _estilizar_header(ws, ["codigo", "descripcion", "categoria"])
    ws.append(["CR-01", "Carcasa con cortes en costado", "Carcasa"])
    ws.append(["CR-02", "Deformación en aro", "Aro"])
    ws.append(["CR-03", "Banda gastada por debajo del límite", "Banda"])
    ws.append(["CR-04", "Daño por impacto (bache)", "Carcasa"])
    ws.append(["CR-05", "Reparación anterior no apta", "Reparación"])
    _agregar_nota(ws, 9, 1, "Códigos únicos de rechazo para clasificar llantas observadas.")

    # ── SHEET 7: RecetasProduccion ─────────────────────────────────
    ws = wb.create_sheet("RecetasProduccion")
    _estilizar_header(ws, [
        "medida_display", "diseno_nombre", "producto_sku", "cantidad", "unidad",
    ])
    ws.append(["205/55 R16", "D-100", "MP-CAUCHO", 2.5, "KG"])
    ws.append(["205/55 R16", "D-100", "MP-CEMENTO", 0.8, "KG"])
    ws.append(["205/55 R16", "D-100", "MP-BANDA", 1.0, "UNidad"])
    _agregar_nota(ws, 7, 1,
                  "Cantidad de cada producto necesaria por llanta. "
                  "producto_sku debe coincidir con hoja Productos.")

    wb.save(str(OUTPUT))
    print(f"[OK] Plantilla creada: {OUTPUT}")


if __name__ == "__main__":
    build()
