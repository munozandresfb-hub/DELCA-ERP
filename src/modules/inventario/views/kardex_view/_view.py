from __future__ import annotations

import os
from datetime import datetime
from decimal import Decimal
from PySide6.QtCore import QDate
from PySide6.QtWidgets import (
    QComboBox,
    QDateEdit,
    QDialog,
    QFileDialog,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

import openpyxl
HAS_OPENPYXL = True

try:
    from fpdf import FPDF
except ImportError:
    FPDF = None  # type: ignore[assignment]

from src.modules.inventario.models.movimiento_inventario_model import (
    MovimientoInventario,
)
from src.modules.inventario.services.producto_service import ProductoService
from src.modules.inventario.views.kardex_view._dialog import _MovimientoFormDialog


class KardexView(QWidget):
    """Kardex — movement ledger with registration and consultation."""

    COLUMNAS = [
        # Columnas del inventario (estado del producto)
        "SKU", "Nombre", "Categoría", "Cantidad UND", "Unidad",
        "Q minima en planta", "Cantidad KG", "Costo Unit.", "Valor Total",
        # Columnas del movimiento
        "Fecha", "Documento de:", "Cant. Und", "Cant. KG",
        "Saldo Und", "Saldo KG", "Referencia", "Observaciones",
    ]

    TIPOS_MOVIMIENTO = ["ENTRADA", "SALIDA", "MERMA", "AJUSTE"]

    def __init__(self) -> None:
        super().__init__()
        self._productos_cache: list[dict] = []
        self._all_movements: list[MovimientoInventario] = []
        # Últimos números de documento usados (se mantienen constantes hasta
        # que el usuario los cambie manualmente — un documento agrupa varios productos)
        self._ultimo_doc_entrada: str | None = None
        self._ultimo_doc_salida: str | None = None
        self.setup_ui()
        self._safe_cargar_inicial()

    def _safe_cargar_inicial(self) -> None:
        try:
            self._cargar_productos()
        except Exception as e:
            print(f"[KardexView] Error cargando productos: {e}")
        try:
            self._cargar_datos()
        except Exception as e:
            print(f"[KardexView] Error cargando datos: {e}")

    # ── UI Setup ────────────────────────────────────────────────────

    def setup_ui(self) -> None:
        layout = QVBoxLayout()

        # Header
        header = QLabel("Kardex / Movimientos de Inventario")
        header.setStyleSheet("font-size: 18px; font-weight: bold; padding: 10px 0;")
        layout.addWidget(header)

        # ── Toolbar (action buttons) ────────────────────────────────
        toolbar = QHBoxLayout()

        btn_ingreso = QPushButton("➕ Ingreso manual")
        btn_ingreso.setStyleSheet(
            "QPushButton { background: #27ae60; color: white; font-weight: bold; "
            "padding: 8px 16px; border-radius: 5px; border: none; }"
        )
        btn_ingreso.clicked.connect(lambda: self._abrir_formulario("ENTRADA"))

        btn_salida = QPushButton("➖ Salida manual")
        btn_salida.setStyleSheet(
            "QPushButton { background: #e74c3c; color: white; font-weight: bold; "
            "padding: 8px 16px; border-radius: 5px; border: none; }"
        )
        btn_salida.clicked.connect(lambda: self._abrir_formulario("SALIDA"))

        btn_ajuste = QPushButton("🔧 Ajuste")
        btn_ajuste.setStyleSheet(
            "QPushButton { background: #f39c12; color: white; font-weight: bold; "
            "padding: 8px 16px; border-radius: 5px; border: none; }"
        )
        btn_ajuste.clicked.connect(lambda: self._abrir_formulario("AJUSTE"))

        btn_documentos = QPushButton("📋 Documentos")
        btn_documentos.setStyleSheet(
            "QPushButton { background: #6c757d; color: white; font-weight: bold; "
            "padding: 8px 16px; border-radius: 5px; border: none; }"
        )
        btn_documentos.clicked.connect(self._abrir_documentos)

        toolbar.addWidget(btn_ingreso)
        toolbar.addWidget(btn_salida)
        toolbar.addWidget(btn_ajuste)
        toolbar.addWidget(btn_documentos)

        # ── Export buttons ───────────────────────────────────────────
        separator = QLabel("  │  ")
        separator.setStyleSheet("color: #ccc; padding: 0 4px;")
        toolbar.addWidget(separator)

        btn_excel = QPushButton("📊 Exportar Excel")
        btn_excel.setStyleSheet(
            "QPushButton { background: #1d6f42; color: white; font-weight: bold; "
            "padding: 8px 16px; border-radius: 5px; border: none; }"
        )
        btn_excel.clicked.connect(self._exportar_excel)

        btn_pdf = QPushButton("📄 Exportar PDF")
        btn_pdf.setStyleSheet(
            "QPushButton { background: #8b0000; color: white; font-weight: bold; "
            "padding: 8px 16px; border-radius: 5px; border: none; }"
        )
        btn_pdf.clicked.connect(self._exportar_pdf)

        toolbar.addWidget(btn_excel)
        toolbar.addWidget(btn_pdf)
        toolbar.addStretch()
        layout.addLayout(toolbar)

        # ── Filters ─────────────────────────────────────────────────
        filter_row = QHBoxLayout()

        filter_row.addWidget(QLabel("Producto:"))
        self.producto_combo = QComboBox()
        self.producto_combo.addItem("Todos", None)
        self.producto_combo.setMinimumWidth(220)
        self.producto_combo.currentIndexChanged.connect(self._cargar_datos)
        filter_row.addWidget(self.producto_combo)

        filter_row.addWidget(QLabel("Documento de:"))
        self.tipo_combo = QComboBox()
        self.tipo_combo.addItems(["Todos", "ENTRADA", "SALIDA", "MERMA", "AJUSTE"])
        self.tipo_combo.currentIndexChanged.connect(self._cargar_datos)
        filter_row.addWidget(self.tipo_combo)

        filter_row.addWidget(QLabel("Desde:"))
        self.fecha_desde = QDateEdit()
        self.fecha_desde.setCalendarPopup(True)
        self.fecha_desde.setDate(QDate.currentDate().addMonths(-1))
        self.fecha_desde.dateChanged.connect(self._cargar_datos)
        filter_row.addWidget(self.fecha_desde)

        filter_row.addWidget(QLabel("Hasta:"))
        self.fecha_hasta = QDateEdit()
        self.fecha_hasta.setCalendarPopup(True)
        self.fecha_hasta.setDate(QDate.currentDate())
        self.fecha_hasta.dateChanged.connect(self._cargar_datos)
        filter_row.addWidget(self.fecha_hasta)

        btn_refresh = QPushButton("🔄")
        btn_refresh.setStyleSheet("font-size: 14px; padding: 4px 12px;")
        btn_refresh.clicked.connect(self._cargar_datos)
        filter_row.addWidget(btn_refresh)

        filter_row.addStretch()
        layout.addLayout(filter_row)

        # ── Table ───────────────────────────────────────────────────
        self.table = QTableWidget()
        self.table.setColumnCount(len(self.COLUMNAS))
        self.table.setHorizontalHeaderLabels(self.COLUMNAS)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)

        layout.addWidget(self.table)
        self.setLayout(layout)

    # ── Data Loading ────────────────────────────────────────────────

    def _cargar_productos(self) -> None:
        productos = ProductoService.listar_productos(solo_activos=True)
        self._productos_cache = []
        self.producto_combo.clear()
        self.producto_combo.addItem("Todos", None)
        for p in productos:
            label = f"{p.nombre} ({p.sku})"
            self.producto_combo.addItem(label, p.id)
            self._productos_cache.append({
                "id": p.id,
                "nombre": p.nombre,
                "sku": p.sku,
                "categoria": p.categoria or "",
                "unidad": p.unidad_medida or "",
                "stock": float(p.stock or 0),
                "stock_kg": float(p.stock_kg or 0),
                "stock_minimo": float(p.stock_minimo or 0),
                "costo": float(p.costo_unitario or 0),
            })

    def _cargar_datos(self) -> None:
        producto_id = self.producto_combo.currentData()
        tipo = self.tipo_combo.currentText()
        desde = self.fecha_desde.date().toPython()
        hasta = self.fecha_hasta.date().toPython()
        desde_dt = datetime(desde.year, desde.month, desde.day, 0, 0, 0)
        hasta_dt = datetime(hasta.year, hasta.month, hasta.day, 23, 59, 59)

        self._all_movements = ProductoService.obtener_kardex_por_filtros(
            producto_id=producto_id,
            tipo=tipo if tipo != "Todos" else None,
            fecha_desde=desde_dt,
            fecha_hasta=hasta_dt,
            limite=500,
        )
        self._render_tabla()

    # ── Table Rendering ─────────────────────────────────────────────

    def _render_tabla(self) -> None:
        movs = self._all_movements
        self.table.setRowCount(len(movs))

        # Running balance (und y kg): process ASC (oldest first)
        saldo = Decimal("0")
        saldo_kg = Decimal("0")
        saldos: dict[int, Decimal] = {}
        saldos_kg: dict[int, Decimal] = {}
        for m in sorted(movs, key=lambda x: x.fecha or datetime.min):
            if m.tipo == "ENTRADA":
                saldo += m.cantidad
                saldo_kg += Decimal(m.cantidad_kg or 0)
            elif m.tipo in ("SALIDA", "MERMA"):
                saldo -= m.cantidad
                saldo_kg -= Decimal(m.cantidad_kg or 0)
            elif m.tipo == "AJUSTE":
                saldo = m.cantidad
                saldo_kg = Decimal(m.cantidad_kg or 0)
            saldos[m.id] = saldo
            saldos_kg[m.id] = saldo_kg

        # Render in DESC order (newest first)
        for row, m in enumerate(movs):
            prod = self._buscar_producto(m.producto_id)

            # ── Columnas del inventario (estado del producto) ──
            sku = prod["sku"] if prod else ""
            nombre = prod["nombre"] if prod else ""
            categoria = prod["categoria"] if prod else ""
            und = prod["stock"] if prod else 0
            unidad = prod["unidad"] if prod else ""
            minimo = prod["stock_minimo"] if prod else 0
            kg = prod["stock_kg"] if prod else 0
            costo = prod["costo"] if prod else 0

            self.table.setItem(row, 0, QTableWidgetItem(sku))
            self.table.setItem(row, 1, QTableWidgetItem(nombre))
            self.table.setItem(row, 2, QTableWidgetItem(categoria))
            self.table.setItem(row, 3, QTableWidgetItem(f"{und:,.2f}"))
            self.table.setItem(row, 4, QTableWidgetItem(unidad))
            self.table.setItem(row, 5, QTableWidgetItem(f"{minimo:,.2f}"))
            self.table.setItem(row, 6, QTableWidgetItem(f"{kg:,.2f}"))
            self.table.setItem(row, 7, QTableWidgetItem(f"${costo:,.2f}" if costo else "$0"))
            self.table.setItem(row, 8, QTableWidgetItem(f"${kg * costo:,.2f}" if costo else "$0"))

            # ── Columnas del movimiento ──
            self.table.setItem(
                row, 9,
                QTableWidgetItem(m.fecha.strftime("%Y-%m-%d %H:%M") if m.fecha else ""),
            )
            self.table.setItem(row, 10, QTableWidgetItem(m.tipo or ""))

            # Cant. Und con signo
            if m.tipo == "ENTRADA":
                cant_str = f"+{m.cantidad}"
            elif m.tipo in ("SALIDA", "MERMA"):
                cant_str = f"-{m.cantidad}"
            else:
                cant_str = str(m.cantidad)
            self.table.setItem(row, 11, QTableWidgetItem(cant_str))

            # Cant. KG con signo
            mkg = Decimal(m.cantidad_kg or 0)
            if m.tipo == "ENTRADA":
                kg_str = f"+{mkg:.2f}"
            elif m.tipo in ("SALIDA", "MERMA"):
                kg_str = f"-{mkg:.2f}"
            else:
                kg_str = f"{mkg:.2f}"
            self.table.setItem(row, 12, QTableWidgetItem(kg_str))

            self.table.setItem(row, 13, QTableWidgetItem(f"{saldos.get(m.id, Decimal('0')):,.2f}"))
            self.table.setItem(row, 14, QTableWidgetItem(f"{saldos_kg.get(m.id, Decimal('0')):,.2f}"))
            self.table.setItem(row, 15, QTableWidgetItem(m.referencia or ""))
            self.table.setItem(row, 16, QTableWidgetItem(m.observaciones or ""))

    def _buscar_producto(self, producto_id: int) -> dict | None:
        for p in self._productos_cache:
            if p["id"] == producto_id:
                return p
        return None

    # ── Export ────────────────────────────────────────────────────────

    def _exportar_excel(self) -> None:
        if not HAS_OPENPYXL:
            QMessageBox.warning(
                self, "Exportar Excel",
                "openpyxl no está instalado.\nEjecute: pip install openpyxl",
            )
            return

        ruta, _ = QFileDialog.getSaveFileName(
            self, "Guardar como Excel", "kardex.xlsx",
            "Excel (*.xlsx)",
        )
        if not ruta:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Kardex"

        # Header row
        header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        header_fill = openpyxl.styles.PatternFill(
            start_color="2c3e50", end_color="2c3e50", fill_type="solid",
        )
        for col, titulo in enumerate(self.COLUMNAS, 1):
            cell = ws.cell(row=1, column=col, value=titulo)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = openpyxl.styles.Alignment(horizontal="center")

        # Running balance (und y kg) para el export
        saldo = Decimal("0")
        saldo_kg = Decimal("0")
        saldos: dict[int, Decimal] = {}
        saldos_kg: dict[int, Decimal] = {}
        for m in sorted(self._all_movements, key=lambda x: x.fecha or datetime.min):
            if m.tipo == "ENTRADA":
                saldo += m.cantidad
                saldo_kg += Decimal(m.cantidad_kg or 0)
            elif m.tipo in ("SALIDA", "MERMA"):
                saldo -= m.cantidad
                saldo_kg -= Decimal(m.cantidad_kg or 0)
            elif m.tipo == "AJUSTE":
                saldo = m.cantidad
                saldo_kg = Decimal(m.cantidad_kg or 0)
            saldos[m.id] = saldo
            saldos_kg[m.id] = saldo_kg

        # Data rows
        for row, m in enumerate(self._all_movements, 2):
            prod = self._buscar_producto(m.producto_id)
            sku = prod["sku"] if prod else ""
            nombre = prod["nombre"] if prod else ""
            categoria = prod["categoria"] if prod else ""
            und = prod["stock"] if prod else 0
            unidad = prod["unidad"] if prod else ""
            minimo = prod["stock_minimo"] if prod else 0
            kg = prod["stock_kg"] if prod else 0
            costo = prod["costo"] if prod else 0
            mkg = float(m.cantidad_kg or 0)

            cant_str = f"+{m.cantidad}" if m.tipo == "ENTRADA" else (
                f"-{m.cantidad}" if m.tipo in ("SALIDA", "MERMA") else str(m.cantidad))
            kg_str = f"+{mkg:.2f}" if m.tipo == "ENTRADA" else (
                f"-{mkg:.2f}" if m.tipo in ("SALIDA", "MERMA") else f"{mkg:.2f}")

            ws.cell(row=row, column=1, value=sku)
            ws.cell(row=row, column=2, value=nombre)
            ws.cell(row=row, column=3, value=categoria)
            ws.cell(row=row, column=4, value=float(und))
            ws.cell(row=row, column=5, value=unidad)
            ws.cell(row=row, column=6, value=float(minimo))
            ws.cell(row=row, column=7, value=float(kg))
            ws.cell(row=row, column=8, value=float(costo))
            ws.cell(row=row, column=9, value=float(kg) * float(costo))
            ws.cell(
                row=row, column=10,
                value=m.fecha.strftime("%Y-%m-%d %H:%M") if m.fecha else "",
            )
            ws.cell(row=row, column=11, value=m.tipo or "")
            ws.cell(row=row, column=12, value=cant_str)
            ws.cell(row=row, column=13, value=kg_str)
            ws.cell(row=row, column=14, value=float(saldos.get(m.id, 0)))
            ws.cell(row=row, column=15, value=float(saldos_kg.get(m.id, 0)))
            ws.cell(row=row, column=16, value=m.referencia or "")
            ws.cell(row=row, column=17, value=m.observaciones or "")

        # Column widths
        widths = [12, 30, 14, 12, 8, 12, 12, 12, 12, 18, 14, 10, 10, 10, 10, 20, 30]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

        wb.save(ruta)
        QMessageBox.information(
            self, "Exportar Excel",
            f"✓ Exportado correctamente:\n{os.path.basename(ruta)}",
        )

    def _exportar_pdf(self) -> None:
        if FPDF is None:
            QMessageBox.warning(
                self, "Exportar PDF",
                "fpdf2 no está instalado.\nEjecute: pip install fpdf2",
            )
            return

        ruta, _ = QFileDialog.getSaveFileName(
            self, "Guardar como PDF", "kardex.pdf",
            "PDF (*.pdf)",
        )
        if not ruta:
            return

        pdf = FPDF(orientation="L", unit="mm", format="A4")
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)

        # Title
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 8, "KARDEX - MOVIMIENTOS DE INVENTARIO", align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        desde_str = self.fecha_desde.date().toString("yyyy-MM-dd")
        hasta_str = self.fecha_hasta.date().toString("yyyy-MM-dd")
        filtros = f"Filtro: {desde_str} al {hasta_str}"
        if self.producto_combo.currentIndex() > 0:
            filtros += f" | Producto: {self.producto_combo.currentText()}"
        if self.tipo_combo.currentIndex() > 0:
            filtros += f" | Documento de: {self.tipo_combo.currentText()}"
        pdf.cell(0, 6, filtros, align="C", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)

        # Column widths (landscape A4 = 297mm, margins ~10mm each)
        col_widths = [11, 24, 13, 9, 6, 9, 9, 10, 10, 15, 11, 8, 8, 8, 8, 17, 25]
        headers = self.COLUMNAS

        # Table header
        pdf.set_font("Helvetica", "B", 7)
        pdf.set_fill_color(44, 62, 80)
        pdf.set_text_color(255, 255, 255)
        for i, h in enumerate(headers):
            pdf.cell(col_widths[i], 6, h, border=1, align="C", fill=True)
        pdf.ln()

        # Table rows
        pdf.set_font("Helvetica", "", 6.5)
        pdf.set_text_color(0, 0, 0)

        # Calculate running balance for PDF
        saldo = Decimal("0")
        saldo_kg = Decimal("0")
        saldos: dict[int, Decimal] = {}
        saldos_kg: dict[int, Decimal] = {}
        for m in sorted(self._all_movements, key=lambda x: x.fecha or datetime.min):
            if m.tipo == "ENTRADA":
                saldo += m.cantidad
                saldo_kg += Decimal(m.cantidad_kg or 0)
            elif m.tipo in ("SALIDA", "MERMA"):
                saldo -= m.cantidad
                saldo_kg -= Decimal(m.cantidad_kg or 0)
            elif m.tipo == "AJUSTE":
                saldo = m.cantidad
                saldo_kg = Decimal(m.cantidad_kg or 0)
            saldos[m.id] = saldo
            saldos_kg[m.id] = saldo_kg

        for m in self._all_movements:
            prod = self._buscar_producto(m.producto_id)
            sku = prod["sku"] if prod else ""
            nombre = prod["nombre"] if prod else ""
            categoria = prod["categoria"] if prod else ""
            und = prod["stock"] if prod else 0
            unidad = prod["unidad"] if prod else ""
            minimo = prod["stock_minimo"] if prod else 0
            kg = prod["stock_kg"] if prod else 0
            costo = prod["costo"] if prod else 0
            mkg = float(m.cantidad_kg or 0)
            cant_str = f"+{m.cantidad}" if m.tipo == "ENTRADA" else (
                f"-{m.cantidad}" if m.tipo in ("SALIDA", "MERMA") else str(m.cantidad))
            kg_str = f"+{mkg:.1f}" if m.tipo == "ENTRADA" else (
                f"-{mkg:.1f}" if m.tipo in ("SALIDA", "MERMA") else f"{mkg:.1f}")

            row_data = [
                sku,
                nombre,
                categoria,
                f"{und:,.0f}",
                unidad,
                f"{minimo:,.0f}",
                f"{kg:,.1f}",
                f"${costo:,.0f}",
                f"${kg * costo:,.0f}",
                m.fecha.strftime("%Y-%m-%d") if m.fecha else "",
                m.tipo or "",
                cant_str,
                kg_str,
                f"{saldos.get(m.id, 0):,.0f}",
                f"{saldos_kg.get(m.id, 0):,.1f}",
                m.referencia or "",
                m.observaciones or "",
            ]

            # Check page break
            if pdf.get_y() + 6 > pdf.page_break_trigger:
                pdf.add_page()
                # Reprint header
                pdf.set_font("Helvetica", "B", 7)
                pdf.set_fill_color(44, 62, 80)
                pdf.set_text_color(255, 255, 255)
                for i, h in enumerate(headers):
                    pdf.cell(col_widths[i], 6, h, border=1, align="C", fill=True)
                pdf.ln()
                pdf.set_font("Helvetica", "", 6.5)
                pdf.set_text_color(0, 0, 0)

            # Row background
            if m.tipo == "ENTRADA":
                pdf.set_fill_color(213, 245, 227)
            elif m.tipo in ("SALIDA", "MERMA"):
                pdf.set_fill_color(250, 219, 216)
            elif m.tipo == "AJUSTE":
                pdf.set_fill_color(252, 243, 207)
            else:
                pdf.set_fill_color(255, 255, 255)

            for i, val in enumerate(row_data):
                pdf.cell(col_widths[i], 6, val, border=1, align="C" if i == 0 else "L", fill=True)
            pdf.ln()

        pdf.output(ruta)
        QMessageBox.information(
            self, "Exportar PDF",
            f"✓ Exportado correctamente:\n{os.path.basename(ruta)}",
        )

    # ── Registration ────────────────────────────────────────────────

    def _abrir_formulario(self, tipo: str) -> None:
        # El número de documento se mantiene constante entre movimientos del
        # mismo tipo (un documento agrupa varios productos) hasta cambiarlo.
        inicial = None
        if tipo == "ENTRADA":
            inicial = self._ultimo_doc_entrada
        elif tipo == "SALIDA":
            inicial = self._ultimo_doc_salida
        dlg = _MovimientoFormDialog(tipo, self, numero_documento_inicial=inicial)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            num = dlg.numero_guardado
            if num:
                if tipo == "ENTRADA":
                    self._ultimo_doc_entrada = num
                elif tipo == "SALIDA":
                    self._ultimo_doc_salida = num
            self._cargar_productos()
            self._cargar_datos()

    def _abrir_documentos(self) -> None:
        """Abre el diálogo de documentos de inventario (ingreso/salida de MP)."""
        from src.modules.inventario.views.kardex_view._dialog import _DocumentosDialog

        dlg = _DocumentosDialog(self)
        dlg.exec()