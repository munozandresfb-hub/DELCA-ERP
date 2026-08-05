"""Utilidades de exportación a Excel."""

import openpyxl
from openpyxl.styles import Font as ExcelFont, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HAS_OPENPYXL = True


class _ExportMixin:
    """Utilidades de exportación de reportes."""

    @staticmethod
    def exportar_excel(
        filename: str,
        headers: list[str],
        rows: list[list[str]],
    ) -> bool:
        """Write a table to an Excel file. Returns True on success."""
        if not HAS_OPENPYXL:
            return False
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Reporte"

            hdr_font = ExcelFont(bold=True, color="FFFFFF")
            hdr_fill = PatternFill(
                start_color="2c3e50", end_color="2c3e50", fill_type="solid",
            )
            for ci, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=ci, value=h)
                cell.font = hdr_font
                cell.fill = hdr_fill
                cell.alignment = Alignment(horizontal="center")

            for ri, row in enumerate(rows, 2):
                for ci, val in enumerate(row, 1):
                    ws.cell(row=ri, column=ci, value=val)

            # Auto-fit column widths (capped)
            for col_idx in range(1, len(headers) + 1):
                max_len = len(str(headers[col_idx - 1]))
                for row_idx in range(2, len(rows) + 2):
                    cell_val = ws.cell(row=row_idx, column=col_idx).value
                    if cell_val:
                        max_len = max(max_len, len(str(cell_val)))
                ws.column_dimensions[
                    get_column_letter(col_idx)
                ].width = min(max_len + 3, 50)

            wb.save(filename)
            return True
        except Exception:
            return False